import os
import ast
import numpy as np
import matplotlib.pyplot as plt
from matplotlib.patches import Circle
from openpyxl import load_workbook


def plot_results_from_excel(
    excel_path,
    dt=0.1,
    goal=None,
    obstacles=None,
    show=True,
    results_dir="Results"
):
    """
    Read simulation results from an Excel workbook and generate:

    1. One XY trajectory plot containing all methods.
    2. Input comparison figure:
           rows = methods
           col 0 = velocity v
           col 1 = angular velocity omega
    3. h comparison figure:
           rows = methods
           col 0 = h_now
           col 1 = h_hmax

    Parameters
    ----------
    excel_path : str
        Path to saved .xlsx results file.

    dt : float
        Simulation timestep.

    goal : array-like or None
        Optional goal position [x_goal, y_goal].

    obstacles : list or None
        Optional list:
            [(cx, cy, radius), ...]

    show : bool
        Show figures interactively.

    results_dir : str
        Directory in which figures are saved.
    """

    obstacles = [] if obstacles is None else obstacles

    # ---------------------------------------------------------
    # Load workbook
    # ---------------------------------------------------------
    wb = load_workbook(
        excel_path,
        data_only=True
    )

    methods = wb.sheetnames
    n_methods = len(methods)

    print(f"Loaded: {excel_path}")
    print(f"Methods found: {methods}")

    all_data = {}

    # =========================================================
    # READ EVERY SHEET
    # =========================================================
    for method in methods:

        ws = wb[method]

        # -----------------------------------------------------
        # Find header row automatically
        #
        # Looks for:
        # step | state_0 | state_1 | ...
        # -----------------------------------------------------
        header_row = None

        for row in range(1, ws.max_row + 1):

            if ws.cell(row=row, column=1).value == "step":
                header_row = row
                break

        if header_row is None:
            print(
                f"Warning: no data header found in sheet '{method}'. "
                "Skipping."
            )
            continue

        # -----------------------------------------------------
        # Read headers
        # -----------------------------------------------------
        headers = [
            ws.cell(row=header_row, column=col).value
            for col in range(1, ws.max_column + 1)
        ]

        # Remove trailing None headers
        while headers and headers[-1] is None:
            headers.pop()

        # -----------------------------------------------------
        # Read numerical data
        # -----------------------------------------------------
        rows = []

        for row in range(header_row + 1, ws.max_row + 1):

            values = [
                ws.cell(row=row, column=col).value
                for col in range(1, len(headers) + 1)
            ]

            # Skip completely empty rows
            if all(v is None for v in values):
                continue

            rows.append(values)

        if len(rows) == 0:
            print(f"Warning: sheet '{method}' contains no data.")
            continue

        # -----------------------------------------------------
        # Convert each column separately
        # -----------------------------------------------------
        column_data = {}

        for j, header in enumerate(headers):

            if header is None:
                continue

            values = []

            for row in rows:

                value = row[j]

                if value is None:
                    values.append(np.nan)
                else:
                    try:
                        values.append(float(value))
                    except (TypeError, ValueError):
                        values.append(np.nan)

            column_data[header] = np.asarray(values)

        all_data[method] = column_data

    # Remove skipped sheets
    methods = list(all_data.keys())
    n_methods = len(methods)

    if n_methods == 0:
        raise ValueError("No valid simulation data found in workbook.")

    # =========================================================
    # COLORS
    # =========================================================
    colors = plt.get_cmap("tab10")(
        np.linspace(0, 1, max(n_methods, 2))
    )

    # =========================================================
    # FIGURE 1: ALL XY TRAJECTORIES
    # =========================================================
    fig_traj, ax_traj = plt.subplots(
        figsize=(8, 8)
    )

    # ---------------------------------------------------------
    # Obstacles
    # ---------------------------------------------------------
    for cx, cy, radius in obstacles:

        circle = Circle(
            (cx, cy),
            radius,
            facecolor="lightcoral",
            edgecolor="darkred",
            alpha=0.5,
            zorder=1
        )

        ax_traj.add_patch(circle)

        ax_traj.plot(
            cx,
            cy,
            "x",
            color="darkred",
            markersize=8,
            zorder=2
        )

    # ---------------------------------------------------------
    # Trajectories
    # ---------------------------------------------------------
    for i, method in enumerate(methods):

        data = all_data[method]

        if "state_0" not in data or "state_1" not in data:
            print(
                f"Warning: state data missing for {method}."
            )
            continue

        x = data["state_0"]
        y = data["state_1"]

        valid = np.isfinite(x) & np.isfinite(y)

        x = x[valid]
        y = y[valid]

        if len(x) == 0:
            continue

        c = colors[i]

        ax_traj.plot(
            x,
            y,
            "-",
            linewidth=1.8,
            color=c,
            label=method,
            zorder=3
        )

        # Initial position
        ax_traj.plot(
            x[0],
            y[0],
            "o",
            color="green",
            markersize=8,
            markeredgecolor="black",
            zorder=5
        )

        # Final position
        ax_traj.plot(
            x[-1],
            y[-1],
            "s",
            color=c,
            markersize=6,
            markeredgecolor="black",
            zorder=5
        )

    # ---------------------------------------------------------
    # Goal
    # ---------------------------------------------------------
    if goal is not None:

        goal = np.asarray(goal).reshape(-1)

        ax_traj.plot(
            goal[0],
            goal[1],
            "*",
            color="gold",
            markersize=16,
            markeredgecolor="black",
            label="Goal",
            zorder=6
        )

    # ---------------------------------------------------------
    # Trajectory formatting
    # ---------------------------------------------------------
    ax_traj.set_xlabel("x [m]")
    ax_traj.set_ylabel("y [m]")
    ax_traj.set_title("Trajectory comparison")

    ax_traj.set_aspect(
        "equal",
        adjustable="datalim"
    )

    ax_traj.grid(
        True,
        alpha=0.3
    )

    ax_traj.legend(
        loc="best",
        fontsize=9
    )

    fig_traj.tight_layout()

    # =========================================================
    # FIGURE 2: INPUTS
    #
    # Row = method
    # Left = v
    # Right = omega
    # =========================================================
    fig_u, axes_u = plt.subplots(
        n_methods,
        2,
        figsize=(12, 3 * n_methods),
        squeeze=False,
        sharex=False
    )

    for i, method in enumerate(methods):

        data = all_data[method]

        c = colors[i]

        ax_v = axes_u[i, 0]
        ax_w = axes_u[i, 1]

        # -----------------------------------------------------
        # Velocity
        # -----------------------------------------------------
        if "input_0" in data:

            v = data["input_0"]

            valid = np.isfinite(v)
            v = v[valid]

            t_v = np.arange(len(v)) * dt

            ax_v.plot(
                t_v,
                v,
                linewidth=1.5,
                color=c,
                drawstyle="steps-post"
            )

        ax_v.set_ylabel(
            f"{method}\n$v$ [m/s]"
        )

        ax_v.grid(
            True,
            alpha=0.3
        )

        # -----------------------------------------------------
        # Angular velocity
        # -----------------------------------------------------
        if "input_1" in data:

            omega = data["input_1"]

            valid = np.isfinite(omega)
            omega = omega[valid]

            t_w = np.arange(len(omega)) * dt

            ax_w.plot(
                t_w,
                omega,
                linewidth=1.5,
                color=c,
                drawstyle="steps-post"
            )

        ax_w.set_ylabel(
            r"$\omega$ [rad/s]"
        )

        ax_w.grid(
            True,
            alpha=0.3
        )

        # -----------------------------------------------------
        # Only bottom row needs time labels
        # -----------------------------------------------------
        if i == n_methods - 1:
            ax_v.set_xlabel("Time [s]")
            ax_w.set_xlabel("Time [s]")

    axes_u[0, 0].set_title(
        "Linear velocity"
    )

    axes_u[0, 1].set_title(
        "Angular velocity"
    )

    fig_u.suptitle(
        "Control input comparison",
        fontsize=14
    )

    fig_u.tight_layout()

    # =========================================================
    # FIGURE 3: h FUNCTIONS
    #
    # Row = method
    # Left = h_now
    # Right = h_hmax
    # =========================================================
    fig_h, axes_h = plt.subplots(
        n_methods,
        2,
        figsize=(12, 3 * n_methods),
        squeeze=False,
        sharex=False
    )

    for i, method in enumerate(methods):

        data = all_data[method]

        c = colors[i]

        ax_now = axes_h[i, 0]
        ax_hmax = axes_h[i, 1]

        # -----------------------------------------------------
        # Find all h_now columns
        #
        # h_now_0
        # h_now_1
        # ...
        # -----------------------------------------------------
        h_now_keys = sorted(
            [
                key
                for key in data.keys()
                if key.startswith("h_now_")
            ]
        )

        # -----------------------------------------------------
        # Find all h_hmax columns
        # -----------------------------------------------------
        h_hmax_keys = sorted(
            [
                key
                for key in data.keys()
                if key.startswith("h_hmax_")
            ]
        )

        # -----------------------------------------------------
        # Plot h_now
        # -----------------------------------------------------
        for j, key in enumerate(h_now_keys):

            h = data[key]

            valid = np.isfinite(h)

            if not np.any(valid):
                continue

            t = np.arange(len(h)) * dt

            ax_now.plot(
                t[valid],
                h[valid],
                linewidth=1.5,
                label=f"h {j}"
            )

        ax_now.axhline(
            0.0,
            color="black",
            linestyle="--",
            linewidth=1.0
        )

        ax_now.set_ylabel(
            f"{method}\nh"
        )

        ax_now.grid(
            True,
            alpha=0.3
        )

        if len(h_now_keys) > 1:
            ax_now.legend(
                fontsize=7,
                loc="best"
            )

        # -----------------------------------------------------
        # Plot h_hmax
        # -----------------------------------------------------
        for j, key in enumerate(h_hmax_keys):

            h = data[key]

            valid = np.isfinite(h)

            if not np.any(valid):
                continue

            t = np.arange(len(h)) * dt

            ax_hmax.plot(
                t[valid],
                h[valid],
                linewidth=1.5,
                label=f"h {j}"
            )

        ax_hmax.axhline(
            0.0,
            color="black",
            linestyle="--",
            linewidth=1.0
        )

        ax_hmax.grid(
            True,
            alpha=0.3
        )

        if len(h_hmax_keys) > 1:
            ax_hmax.legend(
                fontsize=7,
                loc="best"
            )

        if i == n_methods - 1:
            ax_now.set_xlabel("Time [s]")
            ax_hmax.set_xlabel("Time [s]")

    axes_h[0, 0].set_title(
        r"$h(x_t)$"
    )

    axes_h[0, 1].set_title(
        r"$h_{hmax}$"
    )

    fig_h.suptitle(
        "Safety certificate comparison",
        fontsize=14
    )

    fig_h.tight_layout()

    # =========================================================
    # SAVE
    # =========================================================
    os.makedirs(
        results_dir,
        exist_ok=True
    )

    trajectory_path = os.path.join(
        results_dir,
        "excel_trajectory_comparison.png"
    )

    input_path = os.path.join(
        results_dir,
        "excel_input_comparison.png"
    )

    h_path = os.path.join(
        results_dir,
        "excel_h_comparison.png"
    )

    fig_traj.savefig(
        trajectory_path,
        dpi=300,
        bbox_inches="tight"
    )

    fig_u.savefig(
        input_path,
        dpi=300,
        bbox_inches="tight"
    )

    fig_h.savefig(
        h_path,
        dpi=300,
        bbox_inches="tight"
    )

    print(f"Saved trajectory plot: {trajectory_path}")
    print(f"Saved input plot     : {input_path}")
    print(f"Saved h plot         : {h_path}")

    if show:
        plt.show()
    else:
        plt.close(fig_traj)
        plt.close(fig_u)
        plt.close(fig_h)

    return fig_traj, fig_u, fig_h

if __name__ == "__main__":

    plot_results_from_excel(
        excel_path="simulation_20260828_002746.xlsx",
        dt=0.1,
        goal=np.array([4.5, 4.5]),
        obstacles=[
            (2.0, 2.5, 0.5)
        ],
        show=True
    )