"""
Tabulate the disturbance-sweep results written by data_gen_sweep.py.

    python analyze_sweep.py                       # uses ROOTS below
    python analyze_sweep.py sweep_results sweep_results_rpcbf --out sweep_summary.xlsx

Reads every run_*.xlsx under each root, computes per-run metrics, applies the two
start-point filters, and writes one workbook:

    runs               one row per run (all metrics, plus the filter flags)
    discarded          runs removed by the filters and why
    summary            per (sweep, method, controller, env_noise)
    summary_by_radius  same, additionally split by start radius
    paired             per (sweep, env_noise, method/controller): outcome vs. the benchmark
                       on the SAME start point and SAME disturbance realisation

Sign convention (from Main._cbf_terms):  h <= 0 is SAFE, the constraint is hdot <= -alpha*h.
A start point is outside the safe set when the certificate V^h(x0) = max_t h > 0.
"""
import re
import sys
import glob
import argparse
import numpy as np
import pandas as pd
from pathlib import Path
from openpyxl import load_workbook

# ----------------------------------------------------------------------------- configuration
ROOTS = ["sweep_results", "sweep_results_rpcbf"]
OUT = "sweep_summary.xlsx"

OBSTACLES = {                      # must mirror Get_obstacles.ObsDyn layouts: (cx, cy, R)
    "single": [(2.0, 2.5, 0.3)],
    "multi":  [(2.0, 2.5, 0.3), (3.0, 3.5, 0.2), (1.5, 1.8, 0.1)],
}
OBS_MARGIN   = 0.10                # start points closer than R + margin to an obstacle centre are discarded
H_SAFE_MAX   = 0.0                 # safe set is {h_max <= H_SAFE_MAX}
GOAL_TOL     = 0.05                # run_sim stops when dist < 0.05  -> "reached"
SLACK_TOL    = 1e-6                # delta > tol counts as CLF relaxed
CBF_ACT_TOL  = 1e-6                # |cbf_margin| < tol counts as CBF active
BENCHMARK    = {"None", "pure_backup"}   # methods treated as the benchmark inside a sweep
OPTIONAL_COLS = ["h_max_start", "h_max_peak", "frac_h_unsafe", "h_now_peak", "frac_cbf_active",
                 "min_cbf_margin", "frac_clf_relaxed", "mean_slack", "max_slack",
                 "Vmax_start", "Vmax_end", "frac_cert_valid"]


# ----------------------------------------------------------------------------- reading
def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return v


def load_run(path):
    """-> (meta: dict, df: DataFrame of per-step series)"""
    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(rows) if r and r[0] == "step")
    meta = {r[0]: r[1] for r in rows[:hdr_i] if r and r[0] not in (None, "", "Simulation Settings")}
    hdr = [h for h in rows[hdr_i] if h is not None]
    data = [list(r[:len(hdr)]) for r in rows[hdr_i + 1:] if r and r[0] is not None]
    df = pd.DataFrame(data, columns=hdr).apply(pd.to_numeric, errors="coerce")
    return meta, df


def _cols(df, prefix):
    cs = [c for c in df.columns if re.fullmatch(rf"{re.escape(prefix)}_\d+", c)]
    return df[cs].to_numpy(dtype=float) if cs else None


def _parse_list(s):
    return np.array([float(x) for x in re.findall(r"[-+]?\d*\.?\d+(?:e[-+]?\d+)?", str(s))])


# ----------------------------------------------------------------------------- per-run metrics
def run_metrics(meta, df, sweep):
    dt = float(meta.get("dt", 0.01))
    early_stop = int(float(meta.get("early_stop", np.inf)))
    goal = _parse_list(meta.get("init_goal", "[4,1]"))[:2]
    x0 = _parse_list(meta.get("x_s"))
    obs = OBSTACLES.get(str(meta.get("no_obs", "single")), [])

    X = _cols(df, "state")                       # (N, 3)
    U = _cols(df, "input")                       # (N, 2)
    n_steps = len(df)
    pos = X[:, :2]
    seg = np.linalg.norm(np.diff(pos, axis=0), axis=1)
    path_len = float(seg.sum())
    straight = float(np.linalg.norm(goal - x0[:2]))
    final_dist = float(np.linalg.norm(goal - pos[-1]))

    # clearance to obstacles along the path (positive = outside)
    if obs:
        clear = np.min([np.linalg.norm(pos - np.array([cx, cy]), axis=1) - R for cx, cy, R in obs], axis=0)
        min_clear = float(clear.min())
        start_clear = float(clear[0])
    else:
        min_clear = start_clear = np.nan

    col = str(meta.get("collision", "False")).lower() == "true"
    reached = (not col) and final_dist < GOAL_TOL
    timeout = (not col) and (not reached) and n_steps >= early_stop
    infeasible = (not col) and (not reached) and (not timeout)   # QP stopped early

    out = dict(
        sweep=sweep, method=meta.get("method"), controller=meta.get("controller"),
        env_noise=meta.get("env_noise"), rollout_noise=meta.get("rollout_noise"),
        run_id=int(float(meta.get("run_id"))), circle=int(float(meta.get("circle", -1))),
        radius=float(meta.get("radius", np.nan)), point=int(float(meta.get("point", -1))),
        x0=x0[0], y0=x0[1], theta0=x0[2],
        d_scale=str(meta.get("d_scale", "")),
        n_steps=n_steps, T_end=n_steps * dt,
        outcome="collision" if col else "reached" if reached else "timeout" if timeout else "infeasible",
        collision=col, reached=reached, timeout=timeout, infeasible=infeasible,
        collision_t=_num(meta.get("collision_t")) if col else np.nan,
        time_to_goal=n_steps * dt if reached else np.nan,
        final_dist=final_dist,
        path_len=path_len, straight_dist=straight,
        path_eff=straight / path_len if path_len > 0 else np.nan,
        start_clearance=start_clear, min_clearance=min_clear,
        mean_v=float(np.nanmean(U[:, 0])), mean_abs_om=float(np.nanmean(np.abs(U[:, 1]))),
        ctrl_effort=float(np.nansum(np.abs(U[:, 1])) * dt),          # integral |omega| dt
        input_roughness=float(np.nanmean(np.linalg.norm(np.diff(U, axis=0), axis=1))) if len(U) > 1 else np.nan,
    )

    # certificate / barrier statistics (only where the method logs them)
    H = _cols(df, "h_hmax")
    if H is not None and H.size and np.isfinite(H).any():
        out["h_max_start"] = float(np.nanmax(H[0]))                 # V^h(x0), max over obstacles
        out["h_max_peak"] = float(np.nanmax(H))                     # worst certificate value over the run
        out["frac_h_unsafe"] = float(np.mean(np.nanmax(H, axis=1) > H_SAFE_MAX))
    Hn = _cols(df, "h_now")
    if Hn is not None and Hn.size and np.isfinite(Hn).any():
        out["h_now_peak"] = float(np.nanmax(Hn))
    M = _cols(df, "cbf_margin")
    if M is not None and M.size and np.isfinite(M).any():
        out["frac_cbf_active"] = float(np.mean(np.nanmin(np.abs(M), axis=1) < CBF_ACT_TOL))
        out["min_cbf_margin"] = float(np.nanmin(M))
    D = _cols(df, "delta")
    if D is not None and D.size and np.isfinite(D).any():
        d = D[:, 0]
        out["frac_clf_relaxed"] = float(np.nanmean(d > SLACK_TOL))
        out["mean_slack"] = float(np.nanmean(d))
        out["max_slack"] = float(np.nanmax(d))
    V = _cols(df, "V_max")
    if V is not None and V.size and np.isfinite(V).any():
        out["Vmax_start"] = float(V[0, 0]); out["Vmax_end"] = float(V[-1, 0])
    C = _cols(df, "cert_valid")
    if C is not None and C.size and np.isfinite(C).any():
        out["frac_cert_valid"] = float(np.nanmean(C[:, 0]))
    return out


# ----------------------------------------------------------------------------- filters
def apply_filters(runs):
    """Adds discard flags. Filters act on START POINTS (env_noise, run_id) so pairing is preserved."""
    runs = runs.copy()
    for c in OPTIONAL_COLS:                      # methods without a CBF/CLF never log these
        if c not in runs:
            runs[c] = np.nan

    # (1) start point inside / too close to an obstacle -> geometric, same for every batch
    runs["bad_obstacle"] = runs["start_clearance"] < OBS_MARGIN

    # (2) start outside the safe set: V^h(x0) > H_SAFE_MAX. Only CBF methods log h_hmax, so
    #     build the set of bad (env_noise, run_id) from any run that has it, and apply it to all.
    ref = runs.dropna(subset=["h_max_start"])
    bad_h = set(map(tuple, ref.loc[ref["h_max_start"] > H_SAFE_MAX, ["env_noise", "run_id"]].to_numpy()))
    runs["bad_hmax"] = [((n, r) in bad_h) for n, r in zip(runs["env_noise"], runs["run_id"])]
    runs["hmax_reference_available"] = len(ref) > 0

    runs["discard"] = runs["bad_obstacle"] | runs["bad_hmax"]
    runs["discard_reason"] = np.select(
        [runs["bad_obstacle"] & runs["bad_hmax"], runs["bad_obstacle"], runs["bad_hmax"]],
        ["obstacle+hmax", "start inside obstacle (+margin)", "h_max(x0) > 0 (outside safe set)"], "")
    return runs


# ----------------------------------------------------------------------------- aggregation
def summarise(valid, keys):
    rows = []
    for key, g in valid.groupby(keys, dropna=False):
        key = key if isinstance(key, tuple) else (key,)
        r = dict(zip(keys, key))
        r.update(n_runs=len(g),
                 collision_pct=100 * g["collision"].mean(),
                 reached_pct=100 * g["reached"].mean(),
                 timeout_pct=100 * g["timeout"].mean(),
                 infeasible_pct=100 * g["infeasible"].mean(),
                 time_to_goal_mean=g["time_to_goal"].mean(),
                 time_to_goal_median=g["time_to_goal"].median(),
                 path_eff_mean=g["path_eff"].mean(),
                 min_clearance_mean=g["min_clearance"].mean(),
                 min_clearance_worst=g["min_clearance"].min(),
                 ctrl_effort_mean=g["ctrl_effort"].mean(),
                 mean_abs_om=g["mean_abs_om"].mean(),
                 input_roughness_mean=g["input_roughness"].mean())
        # certificate statistics: mean over runs, except peaks/maxima which take the worst run
        for c in ("frac_clf_relaxed", "mean_slack", "frac_cbf_active", "min_cbf_margin",
                  "frac_h_unsafe", "frac_cert_valid"):
            r[c + "_mean"] = g[c].mean() if g[c].notna().any() else np.nan
        for c in ("max_slack", "h_max_peak"):
            r[c + "_worst"] = g[c].max() if g[c].notna().any() else np.nan
        r["min_cbf_margin_worst"] = g["min_cbf_margin"].min() if g["min_cbf_margin"].notna().any() else np.nan
        rows.append(r)
    return pd.DataFrame(rows)


def paired(valid):
    """Outcome of each (method, controller) vs the benchmark on the same (sweep, env_noise, run_id)."""
    rows = []
    for (sweep, noise), g in valid.groupby(["sweep", "env_noise"]):
        bench = g[g["method"].isin(BENCHMARK)]
        if bench.empty:
            continue
        bench = bench.set_index("run_id")
        for (m, c), gm in g[~g["method"].isin(BENCHMARK)].groupby(["method", "controller"]):
            gm = gm.set_index("run_id")
            common = gm.index.intersection(bench.index)
            b, x = bench.loc[common], gm.loc[common]
            rows.append(dict(
                sweep=sweep, env_noise=noise, method=m, controller=c,
                benchmark=str(bench["method"].iloc[0]) + "/" + str(bench["controller"].iloc[0]),
                n_pairs=len(common),
                bench_collision_pct=100 * b["collision"].mean(),
                method_collision_pct=100 * x["collision"].mean(),
                saved_by_method=int((b["collision"] & ~x["collision"]).sum()),
                broken_by_method=int((~b["collision"] & x["collision"]).sum()),
                both_collide=int((b["collision"] & x["collision"]).sum()),
                both_safe=int((~b["collision"] & ~x["collision"]).sum()),
                bench_reached_pct=100 * b["reached"].mean(),
                method_reached_pct=100 * x["reached"].mean(),
                dtime_to_goal_mean=(x["time_to_goal"] - b["time_to_goal"]).mean(),   # + slower than bench
                dpath_len_mean=(x["path_len"] - b["path_len"]).mean(),
                dmin_clearance_mean=(x["min_clearance"] - b["min_clearance"]).mean(),
            ))
    return pd.DataFrame(rows)


# ----------------------------------------------------------------------------- main
def main(roots, out):
    recs = []
    for root in roots:
        files = sorted(glob.glob(str(Path(root) / "*" / "run_*.xlsx")))
        print(f"{root}: {len(files)} runs")
        for f in files:
            try:
                meta, df = load_run(f)
                r = run_metrics(meta, df, sweep=Path(root).name)
                r["file"] = f
                recs.append(r)
            except Exception as e:
                print(f"  skip {f}: {e}")
    if not recs:
        sys.exit("no runs found")

    runs = apply_filters(pd.DataFrame(recs))
    valid = runs[~runs["discard"]]
    print(f"{len(runs)} runs, {int(runs['discard'].sum())} discarded, {len(valid)} kept")
    if not runs["hmax_reference_available"].iloc[0]:
        print("WARNING: no batch logs h_hmax; the h_max(x0) filter could not be applied")

    summary = summarise(valid, ["sweep", "method", "controller", "env_noise"])
    by_radius = summarise(valid, ["sweep", "method", "controller", "env_noise", "radius"])
    pairs = paired(valid)

    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        summary.round(4).to_excel(xw, sheet_name="summary", index=False)
        by_radius.round(4).to_excel(xw, sheet_name="summary_by_radius", index=False)
        pairs.round(4).to_excel(xw, sheet_name="paired", index=False)
        runs.drop(columns=["hmax_reference_available"]).round(6).to_excel(xw, sheet_name="runs", index=False)
        runs[runs["discard"]][["sweep", "method", "controller", "env_noise", "run_id", "radius", "point",
                               "x0", "y0", "start_clearance", "h_max_start", "discard_reason"]] \
            .to_excel(xw, sheet_name="discarded", index=False)
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = max(10, min(28, len(str(col[0].value)) + 2))
    print(f"wrote {out}")
    with pd.option_context("display.width", 200, "display.max_columns", 12):
        print(summary[["sweep", "method", "controller", "env_noise", "n_runs", "collision_pct",
                       "reached_pct", "timeout_pct", "infeasible_pct", "time_to_goal_mean"]].round(2).to_string(index=False))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("roots", nargs="*", default=ROOTS)
    ap.add_argument("--out", default=OUT)
    a = ap.parse_args()
    main(a.roots, a.out)
