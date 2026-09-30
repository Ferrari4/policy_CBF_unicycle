"""
Excel writer for the disturbance sweep (copy of plotter/Save_results.py adapted for batch runs).

Differences from save_results_to_excel:
  * writes into a caller-supplied folder (one per batch) instead of ./logs
  * deterministic file name  run_XXX_rR_pP.xlsx  (no timestamp) so files sort by run_id
  * filter parameters (gamma, alpha, d_scale, ...) come from a plain dict extracted in the
    worker, so the policy_filter object never has to be pickled back through the pool
  * optional job metadata (circle / radius / point) is written into the settings block
"""
import numpy as np
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font

FILTER_PARAM_NAMES = ["gamma", "alpha", "slack_weight", "v_min", "v_max", "om_max",
                      "T_rollout", "T_rollout_clf", "T_dstb_hold", "dt", "n_samples",
                      "n_samples_uniform", "d_scale", "clf_exact_tail", "cbf_delta"]


def extract_filter_params(pf):
    """Pull the outcome-relevant parameters off a policy_filter into a plain dict (call in the worker)."""
    if pf is None:
        return {}
    params = {name: str(getattr(pf, name)) for name in FILTER_PARAM_NAMES if hasattr(pf, name)}
    if hasattr(pf, "clf") and hasattr(pf.clf, "k"):
        params["k (CLF heading weight)"] = str(pf.clf.k)
    return params


def to_2d(arr):
    arr = np.asarray(arr)
    if arr.size == 0:
        return np.empty((0, 1))
    return arr.reshape(len(arr), -1)


def save_sweep_result(results, settings, out_dir, run_id, job_meta=None, filter_params=None):
    """
    results       : dict returned by run_simulation (with or without the "safety" entry)
    settings      : the kwargs passed to run_simulation
    out_dir       : batch folder, created if missing
    run_id        : 0..N-1 within the batch
    job_meta      : optional dict, e.g. {"circle": 2, "radius": 2.0, "point": 7}
    filter_params : dict from extract_filter_params(); if None and results has "safety", extracted here
    """
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    job_meta = job_meta or {}
    tag = ""
    if "radius" in job_meta and "point" in job_meta:
        tag = f"_r{job_meta['radius']:g}_p{int(job_meta['point']):02d}"
    filename = out_dir / f"run_{run_id:03d}{tag}.xlsx"

    if filter_params is None:
        filter_params = extract_filter_params(results.get("safety"))

    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(title=str(settings.get("controller", "run"))[:31])

    # ------------------------------------------------- settings block
    metadata = [["Simulation Settings", ""],
                ["Timestamp", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
                ["run_id", run_id]]
    metadata += [[k, str(v)] for k, v in job_meta.items()]
    metadata += [[k, str(v)] for k, v in settings.items()]
    metadata += [[k, v] for k, v in filter_params.items()]

    # outcome summary (collision record comes from run_simulation when det_collison=True)
    col = results.get("collision")
    states = np.asarray(results["states"])
    goal = np.asarray(settings.get("init_goal", [np.nan, np.nan]), dtype=float)
    metadata += [
        ["n_steps", len(results["inputs"])],
        ["collision", bool(col)],
        ["collision_step", col["step"] if col else ""],
        ["collision_t", col["t"] if col else ""],
        ["collision_obstacle", col["obstacle"] if col else ""],
        ["final_dist_to_goal", float(np.linalg.norm(states[-1, :2] - goal)) if len(states) else ""],
    ]
    for row in metadata:
        ws.append(row)
    ws["A1"].font = Font(bold=True)
    ws.append([])

    # ------------------------------------------------- data block
    series = {
        "state":      results["states"],
        "input":      results["inputs"],
        "d_env":      results.get("d_env",      []),   # applied disturbance per step, if logged
        "h_now":      results.get("h_now",      []),
        "h_hmax":     results.get("h_hmax",     []),
        "V":          results.get("V",          []),
        "V_max":      results.get("V_max",      []),
        "delta":      results.get("delta",      []),
        "Vdot":       results.get("Vdot",       []),
        "alV":        results.get("alV",        []),
        "hdot":       results.get("hdot",       []),
        "alH":        results.get("alH",        []),
        "cbf_margin": results.get("cbf_margin", []),
        "ell0":       results.get("ell0",       []),
        "ellT":       results.get("ellT",       []),
        "cert_valid": results.get("cert_valid", []),
    }
    series = {name: to_2d(arr) for name, arr in series.items()}
    N = min(len(series["state"]), len(series["input"]))

    headers = ["step"]
    for name, arr in series.items():
        headers += [f"{name}_{i}" for i in range(arr.shape[1])]
    ws.append(headers)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True)

    for k in range(N):
        row = [k]
        for arr in series.values():
            row += arr[k].tolist() if k < len(arr) else [None] * arr.shape[1]
        ws.append(row)

    ws.freeze_panes = f"A{header_row + 1}"
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 22

    wb.save(filename)
    return filename
