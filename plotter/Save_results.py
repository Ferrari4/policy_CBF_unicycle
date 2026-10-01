import numpy as np
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font
from pathlib import Path


def to_2d(arr):
    arr = np.asarray(arr)
    if arr.size == 0:
        return np.empty((0, 1))
    return arr.reshape(len(arr), -1)


def _cell(v):
    """Excel-safe scalar: sequences -> 'a,b,c', numpy scalars -> python, bool/None kept."""
    if isinstance(v, (list, tuple, np.ndarray)):
        return ",".join(f"{float(x):.6g}" for x in np.asarray(v).ravel())
    if isinstance(v, np.generic):
        return v.item()
    if isinstance(v, dict):
        return str(v)
    return v


# metric columns in the order they appear on the summary sheet (anything else is appended after)
METRIC_ORDER = ["J_u", "J_nom", "effort_increase", "J_int", "T_int", "frac", "peak",
                "J_u_active", "share_active", "T_run", "n_steps", "reached", "collision",
                "h_closest", "h_min", "n_infeasible"]

# settings copied onto each summary row so a row is self-describing
ROW_SETTINGS = ["method", "controller", "h_controller", "v_controller", "var_slack",
                "rollout_noise_cbf", "rollout_noise_clf", "env_noise", "no_obs", "obs_static",
                "T_rollout", "T_rollout_clf"]


def save_results_to_excel(results_dict, settings, run_id=None, filename=None):
    """
    results_dict : {sheet_name: results}  (results = dict returned by run_simulation)
    settings     : the kwargs the run(s) were launched with. May be a single dict (applies to all)
                   or {sheet_name: dict} when each run has its own settings (sweeps).
    run_id       : optional int suffix for the auto-generated filename
    filename     : explicit path; overrides the auto name (used by run_experiments.py)

    Workbook:
      summary   one row per run: run settings + metrics
      params    key/value of results["params"] from the first run (shared settings)
      <run>     metadata block, blank row, header row, per-step data (the original layout)
    Returns the path written.
    """
    if filename is None:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        suffix = "" if run_id is None else f"_run_{run_id:02d}"
        output_dir = Path("logs")
        output_dir.mkdir(parents=True, exist_ok=True)
        filename = output_dir / f"simulation_{timestamp}{suffix}.xlsx"
    else:
        filename = Path(filename)
        filename.parent.mkdir(parents=True, exist_ok=True)

    per_run_settings = isinstance(settings, dict) and all(k in settings for k in results_dict)
    def settings_for(name):
        return settings[name] if per_run_settings else settings

    wb = Workbook()
    wb.remove(wb.active)

    # -------------------------------------------------
    # summary sheet: one row per run
    # -------------------------------------------------
    ws_sum = wb.create_sheet(title="summary")
    metric_keys = []
    for res in results_dict.values():
        for k in res.get("metrics", {}):
            if k not in metric_keys:
                metric_keys.append(k)
    metric_keys = [k for k in METRIC_ORDER if k in metric_keys] + [k for k in metric_keys if k not in METRIC_ORDER]
    sum_headers = ["run_id"] + ROW_SETTINGS + metric_keys
    ws_sum.append(sum_headers)
    for cell in ws_sum[1]:
        cell.font = Font(bold=True)
    for name, res in results_dict.items():
        st = settings_for(name)
        row = [name]
        row += [_cell(st.get(k, None)) for k in ROW_SETTINGS]
        row += [_cell(res.get("metrics", {}).get(k, None)) for k in metric_keys]
        ws_sum.append(row)
    ws_sum.freeze_panes = "B2"

    # -------------------------------------------------
    # params sheet: key/value from the first run
    # -------------------------------------------------
    first = next(iter(results_dict.values()))
    params = first.get("params")
    if params is None and first.get("safety") is not None and hasattr(first["safety"], "params"):
        params = first["safety"].params()
    if params:
        ws_par = wb.create_sheet(title="params")
        ws_par.append(["parameter", "value"])
        ws_par["A1"].font = Font(bold=True); ws_par["B1"].font = Font(bold=True)
        for k, v in params.items():
            ws_par.append([k, _cell(v)])
        ws_par.column_dimensions["A"].width = 22
        ws_par.column_dimensions["B"].width = 40

    # -------------------------------------------------
    # one sheet per run (original layout, extended)
    # -------------------------------------------------
    for policy_name, results in results_dict.items():
        ws = wb.create_sheet(title=str(policy_name)[:31])
        st = settings_for(policy_name)

        metadata = [
            ["Simulation Settings", ""],
            ["Timestamp", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        ]
        metadata += [[k, str(_cell(v))] for k, v in st.items()]

        # full parameter record of the filter (falls back to the old hand-picked list)
        pf = results.get("safety")
        prm = results.get("params")
        if prm is None and pf is not None and hasattr(pf, "params"):
            prm = pf.params()
        if prm:
            metadata.append(["Filter Parameters", ""])
            metadata += [[k, str(_cell(v))] for k, v in prm.items()]
        elif pf is not None:
            for name in ["gamma", "alpha", "slack_weight", "v_min", "v_max", "om_max",
                         "T_rollout", "T_rollout_clf", "dt", "n_samples", "d_scale",
                         "clf_exact_tail", "cbf_delta"]:
                if hasattr(pf, name):
                    metadata.append([name, str(getattr(pf, name))])
        if pf is not None and hasattr(pf, "clf") and hasattr(pf.clf, "k"):
            metadata.append(["k (CLF heading weight)", str(pf.clf.k)])

        # scalar results
        met = results.get("metrics")
        if met:
            metadata.append(["Metrics", ""])
            metadata += [[k, str(_cell(v))] for k, v in met.items()]

        for row in metadata:
            ws.append(row)
        ws["A1"].font = Font(bold=True)
        ws.append([])                 # blank row

        # -------------------------------------------------
        # per-step data
        # -------------------------------------------------
        series = {
            "t":        results.get("t",        []),
            "state":    results["states"],
            "input":    results["inputs"],
            "u_nom":    results.get("u_nom",    []),   # reference the filter was measured against
            "effort":   results.get("effort",   []),   # ||u_act/u_max||^2
            "dev_sq":   results.get("dev_sq",   []),   # ||(u_act-u_nom)/u_max||^2
            "interv":   results.get("interv",   []),   # ||u_act-u_nom|| >= inter_input
            "cbf_on":   results.get("cbf_on",   []),   # CBF multiplier > 0 (backup: interv)
            "J_u_run":  results.get("J_u_run",  []),   # running integral of effort
            "J_int_run":results.get("J_int_run",[]),   # running integral of dev_sq on cbf_on steps
            "lam_cbf":  results.get("lam_cbf",  []),   # CBF multipliers per obstacle
            "h_now":    results.get("h_now",    []),
            "h_hmax":   results.get("h_hmax",   []),
            "V":        results.get("V",        []),
            "V_max":    results.get("V_max",    []),
            "delta":    results.get("delta",    []),
            "Vdot":     results.get("Vdot",     []),
            "alV":      results.get("alV",      []),
            "hdot":     results.get("hdot",     []),   # grad_h (f + G u_act) + dV_dt, per obstacle
            "alH":      results.get("alH",      []),   # -alpha * h_hmax, per obstacle
            "cbf_margin": results.get("cbf_margin", []),   # alH - hdot  (>= 0 satisfied)
            "ell0":     results.get("ell0",     []),
            "ellT":     results.get("ellT",     []),
            "cert_valid": results.get("cert_valid", []),
            "d_env":    results.get("d_env",    []),
        }
        series = {name: to_2d(np.asarray(arr, dtype=float)) for name, arr in series.items()}

        # states has N+1 rows, inputs N -> use the shorter
        N = min(len(series["state"]), len(series["input"]))

        headers = ["step"]
        for name, arr in series.items():
            headers += [f"{name}_{i}" for i in range(arr.shape[1])] if arr.shape[1] > 1 else [name]
        ws.append(headers)

        header_row = ws.max_row
        for cell in ws[header_row]:
            cell.font = Font(bold=True)

        for k in range(N):
            row = [k]
            for arr in series.values():
                if k < len(arr):
                    row += [None if np.isnan(v) else float(v) for v in arr[k]]
                else:
                    row += [None] * arr.shape[1]
            ws.append(row)

        ws.freeze_panes = f"A{header_row + 1}"
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 22

    wb.save(filename)
    print(f"Saved results to: {filename}")
    return filename