import numpy as np
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font


def to_2d(arr):
    arr = np.asarray(arr)
    if arr.size == 0:
        return np.empty((0, 1))
    return arr.reshape(len(arr), -1)


def save_results_to_excel(all_results, settings, filename=None):

    if filename is None:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"plotter/simulation_{timestamp}.xlsx"

    wb = Workbook()
    wb.remove(wb.active)          # drop the default sheet

    for policy_name, results in all_results.items():

        ws = wb.create_sheet(title=policy_name[:31])

        # -------------------------------------------------
        # Simulation settings
        # -------------------------------------------------
        metadata = [
            ["Simulation Settings", ""],
            ["Timestamp", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        ]
        metadata += [[k, str(v)] for k, v in settings.items()]

        # filter parameters that change the outcome (from the policy_filter object)
        pf = results.get("safety")
        if pf is not None:
            for name in ["gamma", "alpha", "slack_weight", "v_min", "v_max", "om_max",
                         "T_rollout", "dt", "n_samples", "d_scale"]:
                if hasattr(pf, name):
                    metadata.append([name, str(getattr(pf, name))])
            if hasattr(pf, "clf"):
                metadata.append(["k (CLF heading weight)", str(pf.clf.k)])

        for row in metadata:
            ws.append(row)
        ws["A1"].font = Font(bold=True)

        ws.append([])                 # blank row

        # -------------------------------------------------
        # Data
        # -------------------------------------------------
        series = {
            "state":  results["states"],
            "input":  results["inputs"],
            "h_now":  results.get("h_now",  []),
            "h_hmax": results.get("h_hmax", []),
            "V":      results.get("V",      []),
            "delta":  results.get("delta",  []),
            "Vdot":   results.get("Vdot",   []),
            "alV":    results.get("alV",    []),
        }
        series = {name: to_2d(arr) for name, arr in series.items()}

        # states has N+1 rows, inputs N -> use the shorter
        N = min(len(series["state"]), len(series["input"]))

        headers = ["step"]
        for name, arr in series.items():
            headers += [f"{name}_{i}" for i in range(arr.shape[1])]
        ws.append(headers)

        header_row = ws.max_row
        for cell in ws[header_row]:
            cell.font = Font(bold=True)

        for k in range(N):
            row = [k]
            for arr in series.values():
                row += arr[k].tolist() if k < len(arr) else [None] * arr.shape[1]
            ws.append(row)

        ws.freeze_panes = f"A{header_row + 1}"
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 22

    wb.save(filename)
    print(f"Saved results to: {filename}")